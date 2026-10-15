from copy import deepcopy
from io import BytesIO
from urllib.parse import unquote
from docx import Document
from starlette.testclient import TestClient
from finbridge.real_case import real_cases
from finbridge.web_api import create_app
from finbridge.workflow import DiligenceRun
from finbridge.word_report import report_docx


def test_word_report_is_editable_and_preserves_failed_decision():
    case=real_cases()[2]
    run=DiligenceRun(case['data'],case['profile'],'融资组合','真实算例').local()
    before=deepcopy(run)
    doc=Document(BytesIO(report_docx(run)))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert all(h in text for h in ('一 财务分析','二 融资决策','三 风险评估','四 方案建议与使用边界'))
    assert '本次未形成通过验证的新增融资组合' in text
    assert '预计可用于偿债的现金' in text and '不是贷款本金损失' in text
    assert len(doc.tables)>=4
    assert run==before


def test_word_download_returns_docx_with_chinese_filename(tmp_path):
    with TestClient(create_app(database_path=tmp_path/'word.sqlite')) as c:
        from test_decision_chat import wait_run
        c.get('/api/bootstrap')
        start=c.post('/api/runs',json={'profile':{'requested_amount':3e8},'use_ai':False})
        run=wait_run(c,start.json()['id'])
        response=c.get('/api/runs/'+run['id']+'/download?format=docx')
        assert response.status_code==200
        assert 'wordprocessingml.document' in response.headers['content-type']
        header=unquote(response.headers['content-disposition'])
        assert 'attachment;' in header and '融资分析报告' in header and '.docx' in header
        doc=Document(BytesIO(response.content))
        assert '企业财务分析与融资决策报告'==doc.paragraphs[0].text
        assert not list(doc.styles['Title'].element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr'))
        assert 'wordprocessingml.document' in c.get('/api/runs/'+run['id']+'/download').headers['content-type']
