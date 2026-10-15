from io import BytesIO

import pandas as pd
import pytest

from finbridge.data import FIELD_CATALOG, FINANCIAL_COLUMNS, load_uploaded_financial_data


def test_load_csv_and_fill_missing_financial_columns():
    content = b"company_name,revenue\nSimulated Co,1000\n"

    frame = load_uploaded_financial_data("companies.csv", content)

    assert frame.loc[0, "company_name"] == "Simulated Co"
    assert frame.loc[0, "revenue"] == 1000
    assert all(column in frame.columns for column in FINANCIAL_COLUMNS)
    assert pd.isna(frame.loc[0, "total_assets"])


def test_load_xlsx():
    buffer = BytesIO()
    pd.DataFrame({"company_name": ["Simulated Co"], "revenue": [1000]}).to_excel(
        buffer, index=False
    )

    frame = load_uploaded_financial_data("companies.xlsx", buffer.getvalue())

    assert frame.loc[0, "company_name"] == "Simulated Co"
    assert frame.loc[0, "revenue"] == 1000


def test_load_chinese_financial_headers_from_xlsx():
    buffer = BytesIO()
    pd.DataFrame(
        {
            "公司名称": ["寒武纪"],
            "资产总计（元）": [1000],
            "负债合计": [500],
            "流动资产合计": [600],
            "流动负债合计": [300],
            "营业总收入（元）": [800],
            "净利润（净亏损以“-”号填列）": [80],
            "报告期": ["2025年"],
        }
    ).to_excel(buffer, index=False)

    frame = load_uploaded_financial_data("寒武纪.xlsx", buffer.getvalue())

    assert frame.loc[0, "company_name"] == "寒武纪"
    assert [frame.loc[0, column] for column in FINANCIAL_COLUMNS] == [
        1000,
        500,
        600,
        300,
        800,
        80,
    ]
    assert frame.loc[0, "report_period"] == "2025年"


def test_load_xlsx_with_title_rows_and_ten_thousand_yuan_unit():
    buffer = BytesIO()
    pd.DataFrame(
        {
            "公司名称": ["寒武纪"],
            "资产总计": [100],
            "负债合计": [40],
            "流动资产合计": [60],
            "流动负债合计": [20],
            "营业收入": [80],
            "净利润": [8],
        }
    ).to_excel(buffer, index=False, startrow=2)
    from openpyxl import load_workbook

    workbook = load_workbook(buffer)
    workbook.active["A1"] = "寒武纪 2025年年度报告（单位：万元）"
    titled = BytesIO()
    workbook.save(titled)

    frame = load_uploaded_financial_data("寒武纪.xlsx", titled.getvalue())

    assert frame.loc[0, "report_period"] == "2025年"
    assert frame.loc[0, "total_assets"] == 1_000_000
    assert frame.loc[0, "net_profit"] == 80_000


def test_merge_financial_sheets_only_when_company_and_period_match():
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(
            {
                "公司名称": ["寒武纪"],
                "报告期": ["2025年度"],
                "资产总计": [100],
                "负债合计": [40],
                "流动资产合计": [60],
                "流动负债合计": [20],
            }
        ).to_excel(writer, sheet_name="资产负债表", index=False)
        pd.DataFrame(
            {
                "公司名称": ["寒武纪"],
                "报告期": ["2025年度"],
                "营业收入": [80],
                "净利润": [8],
            }
        ).to_excel(writer, sheet_name="利润表", index=False)

    frame = load_uploaded_financial_data("寒武纪.xlsx", buffer.getvalue())

    assert len(frame) == 1
    assert frame.loc[0, "total_assets"] == 100
    assert frame.loc[0, "revenue"] == 80
    assert frame.loc[0, "source_sheet"] == "利润表、资产负债表"


def test_reject_multiple_headers_for_same_financial_field():
    buffer = BytesIO()
    pd.DataFrame(
        {"company_name": ["Cambricon"], "公司名称": ["寒武纪"]}
    ).to_excel(buffer, index=False)

    with pytest.raises(ValueError, match="多个表头对应同一字段"):
        load_uploaded_financial_data("companies.xlsx", buffer.getvalue())


@pytest.mark.parametrize(
    "filename,content,error",
    [
        ("companies.json", b"{}", "仅支持"),
        ("companies.csv", b"revenue\n1000\n", "company_name"),
        ("companies.csv", b"company_name\n", "没有企业记录"),
    ],
)
def test_reject_invalid_uploads(filename, content, error):
    with pytest.raises(ValueError, match=error):
        load_uploaded_financial_data(filename, content)


def test_catalog_units_keep_metadata_and_channel_rates_out_of_yuan_scaling():
    content = "企业名称,报告期,营业收入（万元）,经营年限（年）,信用贷款期限（月）,信用贷款年利率（%）,应收融资比例（%）\n测试企业,2025年度,100,8,24,5.2,70\n".encode("utf-8-sig")
    row = load_uploaded_financial_data("中文.csv", content).iloc[0]
    assert row["revenue"] == 1_000_000
    assert row["years_in_operation"] == 8
    assert row["credit_months"] == 24
    assert row["credit_rate"] == pytest.approx(.052)
    assert row["receivable_advance_rate"] == pytest.approx(.7)


@pytest.mark.parametrize("header,value,expected", [
    ("credit_rate", "0.052", .052), ("credit_rate", "5.2%", .052),
    ("信用贷款年利率（%）", "5.2", .052), ("信用贷款年利率（%）", "5.2%", .052),
    ("预计营业收入增长率（%）", "-10", -.1),
])
def test_percentage_strings_headers_and_decimal_keys(header, value, expected):
    content = f"company_name,{header}\n模拟企业,{value}\n".encode("utf-8-sig")
    row = load_uploaded_financial_data("rates.csv", content).iloc[0]
    key = "projected_revenue_growth" if header.startswith("预计") else "credit_rate"
    assert row[key] == pytest.approx(expected)


def test_excel_percentage_format_is_not_divided_twice():
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(["企业名称", "信用贷款年利率（%）", "应收融资比例（%）"])
    workbook.active.append(["模拟企业", .052, 70])
    workbook.active["B2"].number_format = "0.0%"
    output = BytesIO()
    workbook.save(output)
    row = load_uploaded_financial_data("percent.xlsx", output.getvalue()).iloc[0]
    assert row["credit_rate"] == pytest.approx(.052)
    assert row["receivable_advance_rate"] == pytest.approx(.7)


def test_profit_and_cash_definitions_are_not_silent_aliases():
    content = "企业名称,营业利润（元）,息税前利润（元）,货币资金（元）,现金及现金等价物期末余额（元）\n模拟企业,100,120,80,60\n".encode("utf-8-sig")
    row = load_uploaded_financial_data("definitions.csv", content).iloc[0]
    assert row["operating_profit"] == 100
    assert row["ebit"] == 120
    assert row["cash_and_equivalents"] == 80
    assert row["cash_equivalents_balance"] == 60


def test_gb18030_chinese_csv_is_readable():
    frame = load_uploaded_financial_data("中文.csv", "企业名称,营业收入（元）\n模拟企业,100\n".encode("gb18030"))
    assert frame.iloc[0]["revenue"] == 100


def test_catalog_describes_true_sources_and_never_uses_profit_as_cfads():
    assert all(set(entry) == {"label", "group", "unit", "source", "meaning", "required"} for entry in FIELD_CATALOG.values())
    assert "净利润" in FIELD_CATALOG["cash_flow_available_for_debt_service"]["meaning"]
    assert "不能" in FIELD_CATALOG["operating_profit"]["meaning"]
