from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import pytest

from finbridge.data import FIELD_CATALOG, load_uploaded_financial_data
from finbridge.template import SIMULATED_EXAMPLE, legacy_financial_template_csv as financial_template_csv, legacy_financial_template_xlsx as financial_template_xlsx


@pytest.mark.parametrize("format", ["xlsx", "csv"])
def test_complete_chinese_example_roundtrips_all_financial_inputs(format):
    content = financial_template_xlsx(True) if format == "xlsx" else financial_template_csv(True)
    frame = load_uploaded_financial_data("示例." + format, content)
    assert len(frame) == 1
    row = frame.iloc[0]
    for key, expected in SIMULATED_EXAMPLE.items():
        if isinstance(expected, (int, float)):
            assert row[key] == pytest.approx(expected), key
        else:
            assert row[key] == expected, key
    assert "模拟" in row["company_name"]
    assert row["scheduled_debt_service"] == row["annual_principal_repayment"] + row["annual_interest_payment"]


@pytest.mark.parametrize("format", ["xlsx", "csv"])
def test_blank_template_is_not_a_fabricated_company_record(format):
    content = financial_template_xlsx() if format == "xlsx" else financial_template_csv()
    with pytest.raises(ValueError, match="没有企业记录"):
        load_uploaded_financial_data("模板." + format, content)


def test_workbook_has_clear_chinese_tabs_editable_inputs_and_unimported_examples():
    workbook = load_workbook(BytesIO(financial_template_xlsx()))
    assert workbook.sheetnames == ["说明", "企业财务数据", "字段说明", "模拟示例"]
    main = workbook["企业财务数据"]
    assert main["A2"].value == "请填写企业名称"
    assert main["B2"].value is None
    assert main.max_column == len(FIELD_CATALOG)
    assert main.freeze_panes == "C2"
    assert len(main.data_validations.dataValidation) > 60
    assert workbook.active.title == "企业财务数据"
    # Filling only the editable table must not import the reference example.
    main["A2"] = "用户提供企业"
    main["B2"] = "2025年度"
    output = BytesIO()
    workbook.save(output)
    frame = load_uploaded_financial_data("已填写.xlsx", output.getvalue())
    assert len(frame) == 1
    assert frame.iloc[0]["company_name"] == "用户提供企业"
    assert pd.isna(frame.iloc[0]["revenue"])


def test_saved_chinese_files_are_available_and_match_function_content():
    root = Path(__file__).resolve().parents[1] / "data"
    for example, name in ((False, "模板"), (True, "示例")):
        assert (root / f"企业融资分析{name}.xlsx").is_file()
        saved=(root / f"企业融资分析{name}.csv").read_bytes()
        assert saved.startswith(b"\xef\xbb\xbf")
        if example:
            row=load_uploaded_financial_data("历史.csv",saved).iloc[0]
            assert row["revenue"]==SIMULATED_EXAMPLE["revenue"]


def test_example_has_multiple_distinct_channels_and_sufficient_asset_capacity():
    example = SIMULATED_EXAMPLE
    gap = example["financing_need"] - example["available_internal_funds"]
    candidates = [example["confirmed_credit_limit"], example["eligible_receivables"] * example["receivable_advance_rate"], example["unencumbered_collateral_value"] * example["collateral_advance_rate"]]
    assert all(value > 0 for value in candidates)
    assert sum(candidates) >= gap
    assert example["cash_flow_available_for_debt_service"] > example["scheduled_debt_service"]
