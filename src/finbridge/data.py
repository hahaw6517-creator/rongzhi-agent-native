"""Sample and uploaded financial data helpers."""

from __future__ import annotations

from io import BytesIO
from numbers import Real
from pathlib import Path
import re
import unicodedata

import pandas as pd


def _field(label: str, group: str, unit: str, source: str, meaning: str, required: bool = False) -> dict:
    return {"label": label, "group": group, "unit": unit, "source": source, "meaning": meaning, "required": required}


# Values use yuan, decimal ratios and months internally. This catalog is shared by
# the Chinese workbook, import mapping and the application's editing controls.
FIELD_CATALOG = {
    "company_name": _field("企业名称", "企业与期间", "", "营业执照", "填写分析主体完整名称；模板占位名称须替换。", True),
    "report_period": _field("报告期", "企业与期间", "", "财务报表", "年度分析建议填写2025年度等完整年度；各报表必须同主体同期间。", True),
    "industry": _field("所属行业", "企业与期间", "", "企业档案", "填写主营业务所属行业，用于核验经营周期，不代表融资资格。"),
    "region": _field("所在地区", "企业与期间", "", "营业执照", "填写省市；地方政策仍需按机构和地区另行核验。"),
    "years_in_operation": _field("经营年限", "企业与期间", "年", "营业执照及经营记录", "截至报告期的实际持续经营年数。"),
    "purpose": _field("融资用途", "企业与期间", "", "资金用途计划", "填写经营周转、订单备货、设备购置或综合融资；综合融资应拆分资金用途。"),
    "total_assets": _field("资产总额", "资产与现金", "元", "资产负债表", "报告期末资产合计。", True),
    "total_liabilities": _field("负债总额", "资产与现金", "元", "资产负债表", "报告期末负债合计。", True),
    "current_assets": _field("流动资产", "资产与现金", "元", "资产负债表", "报告期末流动资产合计。", True),
    "current_liabilities": _field("流动负债", "资产与现金", "元", "资产负债表", "报告期末流动负债合计。", True),
    "quick_assets": _field("速动资产", "资产与现金", "元", "资产负债表及附注", "按企业核实口径排除存货、预付款等变现受限项目；不能只凭名称估算。"),
    "cash_and_equivalents": _field("货币资金", "资产与现金", "元", "资产负债表及货币资金附注", "为兼容原项目保留该英文键。按资产负债表货币资金填报，受限部分单列；不得无条件使用现金流量表现金等价物余额替代。"),
    "cash_equivalents_balance": _field("现金及现金等价物余额", "资产与现金", "元", "现金流量表补充资料", "现金流量表期末现金及现金等价物余额；与货币资金分别保存，不自动互相替代。"),
    "restricted_cash": _field("受限货币资金", "资产与现金", "元", "货币资金附注及银行证明", "货币资金中保证金、冻结款等受限部分；不是额外资金，不得再次加总。"),
    "beginning_total_assets": _field("期初总资产", "资产与现金", "元", "资产负债表", "与期末总资产同口径，用于平均资产计算。"),
    "ending_total_assets": _field("期末总资产", "资产与现金", "元", "资产负债表", "应与资产总额保持一致。"),
    "beginning_equity": _field("期初所有者权益", "资产与现金", "元", "资产负债表", "期初所有者权益合计。"),
    "ending_equity": _field("期末所有者权益", "资产与现金", "元", "资产负债表", "期末所有者权益合计，应与资产减负债的口径一致。"),
    "revenue": _field("营业收入", "利润与增长", "元", "利润表", "完整报告期营业收入，不能以开票额或回款额直接替代。", True),
    "net_profit": _field("净利润", "利润与增长", "元", "利润表", "本报告期净利润，可为负数；不是可用于偿债现金流。", True),
    "gross_profit": _field("毛利", "利润与增长", "元", "利润表及成本明细", "营业收入减对应营业成本；不是毛利率。"),
    "cost_of_goods_sold": _field("营业成本", "利润与增长", "元", "利润表", "与营业收入对应的营业成本。"),
    "operating_profit": _field("营业利润", "利润与增长", "元", "利润表", "按利润表原值填报；不能无条件等同息税前利润。"),
    "ebit": _field("息税前利润", "利润与增长", "元", "利润表及经核实调整明细", "EBIT须明确非经营损益等调整口径，不将营业利润自动当作EBIT。"),
    "interest_expense": _field("利息费用", "利润与增长", "元", "财务费用附注", "权责发生制利息费用；与实际支付利息分别填报。"),
    "income_tax_expense": _field("所得税费用", "利润与增长", "元", "利润表及所得税附注", "本期所得税费用，包含适用的递延影响；不等同实缴税费。"),
    "depreciation_amortization": _field("折旧及摊销", "利润与增长", "元", "现金流量表补充资料及附注", "本报告期折旧及摊销合计，属于非现金费用。"),
    "prior_revenue": _field("上期营业收入", "利润与增长", "元", "上年同期利润表", "使用与本期长度一致的上年同期口径。"),
    "prior_net_profit": _field("上期净利润", "利润与增长", "元", "上年同期利润表", "使用与本期长度一致的上年同期净利润，可为负数。"),
    "projected_revenue_growth": _field("预计营业收入增长率", "利润与增长", "%", "经核实经营预算", "经营预算预测假设，不是已实现增速；Excel填10%或0.10，CSV百分比列及网站填10表示10%。"),
    "operating_cash_flow": _field("经营活动现金流量净额", "现金流与偿债", "元", "现金流量表", "本报告期经营活动现金净流量，可为负数。"),
    "cash_flow_available_for_debt_service": _field("可用于偿债现金流", "现金流与偿债", "元", "年度现金流预算及口径调节表", "CFADS应由企业确认并提供调节口径，已反映相关经营税费与维持性支出；不得直接用净利润替代，也不得重复扣税费。", True),
    "maintenance_capex": _field("维持性资本开支", "现金流与偿债", "元", "固定资产计划及现金流预算", "维持现有经营所需资本支出，是总资本开支的一部分。"),
    "capital_expenditure": _field("总资本开支", "现金流与偿债", "元", "现金流量表及投资计划", "购建长期资产等资本支出总额；与维持性资本开支不能重复加总。"),
    "taxes_paid": _field("支付的各项税费", "现金流与偿债", "元", "现金流量表及纳税记录", "实际支付税费，通常已在经营现金流中反映；不是所得税费用。"),
    "scheduled_debt_service": _field("现有年度偿债本息", "现金流与偿债", "元", "现有债务还本付息计划", "未来一年现有债务计划偿还本金加支付利息，不含本次拟新增融资。", True),
    "annual_principal_repayment": _field("年度应偿还本金", "现金流与偿债", "元", "现有借款及租赁合同", "未来一年现有债务应还本金，不等于借款余额。"),
    "annual_interest_payment": _field("年度应支付利息", "现金流与偿债", "元", "现有借款及租赁合同", "未来一年现有债务应支付利息；与本期利息费用区分。"),
    "short_term_debt": _field("短期有息债务", "债务结构", "元", "资产负债表及借款附注", "报告期末短期有息债务余额，不含已另列的一年内到期长期债务；与长期有息债务、一年内到期长期债务三项互斥，避免重复统计。"),
    "long_term_debt": _field("长期有息债务", "债务结构", "元", "资产负债表及借款附注", "一年以后到期的有息债务余额，不含已单列的一年内到期长期债务；与短期有息债务、一年内到期长期债务三项互斥。"),
    "current_portion_long_term_debt": _field("一年内到期长期债务", "债务结构", "元", "资产负债表及借款附注", "长期债务中未来一年到期、已单列的部分，不再计入短期或长期有息债务；三项互斥。该余额不是全年计划偿债本金的自动替代。"),
    "accounts_receivable": _field("应收账款", "营运资金与应收", "元", "资产负债表及应收明细", "期末应收账款；与期末应收账款同口径。"),
    "beginning_accounts_receivable": _field("期初应收账款", "营运资金与应收", "元", "资产负债表", "与期末余额同口径，用于周转测算。"),
    "ending_accounts_receivable": _field("期末应收账款", "营运资金与应收", "元", "资产负债表", "期末应收账款余额，需与应收账款字段核对。"),
    "receivables_overdue": _field("逾期应收账款", "营运资金与应收", "元", "应收账款账龄与到期明细", "已经超过合同付款日的应收余额；账龄长不一定已经逾期。"),
    "receivables_over_90_days": _field("逾期90天以上应收账款", "营运资金与应收", "元", "应收账款到期明细", "逾期超过90天部分，为逾期应收的子集，不能重复加总。"),
    "eligible_receivables": _field("合格可融资应收账款", "营运资金与应收", "元", "合同发票验收确权及登记核验", "人工核验真实、权属清晰、未被重复占用的应收债权；账面余额不等于合格债权。"),
    "inventory": _field("存货", "营运资金与应收", "元", "资产负债表及存货明细", "期末存货；与期末存货字段同口径。"),
    "beginning_inventory": _field("期初存货", "营运资金与应收", "元", "资产负债表", "期初存货余额。"),
    "ending_inventory": _field("期末存货", "营运资金与应收", "元", "资产负债表", "期末存货余额。"),
    "accounts_payable": _field("应付账款", "营运资金与应收", "元", "资产负债表", "期末应付账款余额。"),
    "beginning_accounts_payable": _field("期初应付账款", "营运资金与应收", "元", "资产负债表", "期初应付账款余额。"),
    "ending_accounts_payable": _field("期末应付账款", "营运资金与应收", "元", "资产负债表", "期末应付账款余额，应与应付账款字段核对。"),
    "prepayments": _field("预付账款", "营运资金与应收", "元", "资产负债表", "期末预付账款余额。"),
    "beginning_prepayments": _field("期初预付账款", "营运资金与应收", "元", "资产负债表", "期初预付账款余额。"),
    "ending_prepayments": _field("期末预付账款", "营运资金与应收", "元", "资产负债表", "期末预付账款余额。"),
    "advance_receipts": _field("预收账款", "营运资金与应收", "元", "资产负债表及合同负债附注", "按同一口径填报预收款；合同负债是否纳入须说明，不得重复统计。"),
    "beginning_advance_receipts": _field("期初预收账款", "营运资金与应收", "元", "资产负债表及合同负债附注", "期初预收款，口径须与期末一致。"),
    "ending_advance_receipts": _field("期末预收账款", "营运资金与应收", "元", "资产负债表及合同负债附注", "期末预收款，口径须与期初一致。"),
    "financing_need": _field("资金需求总额", "融资需求与资产", "元", "资金用途预算", "本次资金用途总额，不是扣自有资金后的贷款申请额。", True),
    "available_internal_funds": _field("可用自有资金", "融资需求与资产", "元", "资金计划及可支配现金核验", "本次可实际投入的自有资金，应扣受限部分并保留必要经营现金，不直接等于货币资金。", True),
    "confirmed_additional_internal_funds": _field("已落实额外自有资金", "融资需求与资产", "元", "已到账股东投入及独立资金来源证明", "报表货币资金之外已落实、可投入本次用途的非借款资金；排除账面已有现金和已用于最低留存的资金，不把待回款、未承诺出资或新增贷款计入。未填不假定存在额外来源。"),
    "minimum_cash_reserve": _field("最低经营现金储备", "融资需求与资产", "元", "经营现金预算", "企业应保留的最低经营现金，不属于可供本次采购使用的资金。"),
    "working_capital_need": _field("营运资金需求量", "融资需求与资产", "元", "周转测算与经营预算", "按周转和交易需求核验的营运资金总需求，尚未扣自有资金与已有融资。"),
    "working_capital_funding_need": _field("经营周转资金需求", "融资需求与资产", "元", "本次用途预算", "本次总需求中用于周转的部分；与设备需求相加应与总需求核对。"),
    "equipment_funding_need": _field("设备资金需求", "融资需求与资产", "元", "本次用途预算及采购计划", "本次总需求中用于设备的部分，避免用流贷覆盖设备投资。"),
    "equipment_purchase_amount": _field("设备采购金额", "融资需求与资产", "元", "设备报价及采购合同", "经核验设备采购金额，不等于可出租资产价值或获批租赁金额。"),
    "unencumbered_collateral_value": _field("未占用担保物价值", "融资需求与资产", "元", "权属证明评估与登记核验", "核验权属及既有负担后的可用担保物价值，不是融资额度。"),
    "confirmed_credit_limit": _field("已确认可用信用额度", "融资需求与资产", "元", "金融机构书面授信及剩余额度证明", "本次可使用的有效剩余额度，须核验用途、期限、撤销条件及额度是否被占用。"),
    "existing_working_capital_loans": _field("现有流动资金贷款", "融资需求与资产", "元", "借款合同及余额明细", "营运资金测算时应扣除的现有流贷余额；不是追加资金来源。"),
    "other_working_capital_financing": _field("其他营运资金融资", "融资需求与资产", "元", "其他融资合同及余额明细", "已提供营运资金的其他渠道融资，不能与现有流贷或本次拟融资重复。"),
}

for _prefix, _label in (("credit", "信用贷款"), ("receivable", "应收融资"), ("secured", "抵质押贷款"), ("leasing", "设备融资租赁")):
    FIELD_CATALOG[f"{_prefix}_rate"] = _field(f"{_label}年利率", "融资渠道条件", "%", "机构书面报价或用户测算假设", "统一使用年化利率。Excel填5.2%或0.052，CSV百分比列及网站填5.2表示5.2%；未获报价时是测算假设。")
    FIELD_CATALOG[f"{_prefix}_months"] = _field(f"{_label}期限", "融资渠道条件", "月", "机构条款或用户测算假设", "填写融资期限月数；需与经营回款、应收账期或设备寿命相匹配。")
    FIELD_CATALOG[f"{_prefix}_fee_rate"] = _field(f"{_label}一次性费率", "融资渠道条件", "%", "机构收费明细或用户测算假设", "按本次融资本金计的一次性收费比例，不含利息；真实收费须核验，不将一次性费率当年费率。")
for _key, _label in (("receivable_advance_rate", "应收融资比例"), ("collateral_advance_rate", "担保物融资比例"), ("leasing_advance_rate", "设备租赁融资比例")):
    FIELD_CATALOG[_key] = _field(_label, "融资渠道条件", "%", "机构条款或用户测算假设", "合格资产可融资比例。Excel填70%或0.70，CSV百分比列及网站填70表示70%；不是行业标准或授信承诺。")

from finbridge.public_data import BLOCKS as PUBLIC_BLOCKS
for _block, _rows in PUBLIC_BLOCKS.items():
    for _key, _label, _prior, _signed in _rows:
        if _key not in FIELD_CATALOG:
            FIELD_CATALOG[_key] = _field(_label, _block, "元", "公开"+_block, "公开报告原值；未披露留空。")
        if _prior and _prior not in FIELD_CATALOG:
            FIELD_CATALOG[_prior] = _field("上期"+_label, _block, "元", "公开"+_block, "上期同口径原值；未披露留空。")

FINANCIAL_COLUMNS = (
    "total_assets",
    "total_liabilities",
    "current_assets",
    "current_liabilities",
    "revenue",
    "net_profit",
)

_HEADER_ALIASES = {
    "company_name": ("company_name", "company name", "公司名称", "企业名称", "公司", "企业", "证券简称", "股票简称"),
    "report_period": ("report_period", "report period", "period", "报告期", "会计期间", "报告日期", "截止日期", "年度"),
    "total_assets": ("total_assets", "total assets", "资产总计", "资产总额", "资产合计", "总资产"),
    "total_liabilities": ("total_liabilities", "total liabilities", "负债合计", "负债总计", "负债总额", "总负债"),
    "current_assets": ("current_assets", "current assets", "流动资产合计", "流动资产总计", "流动资产总额", "流动资产"),
    "current_liabilities": ("current_liabilities", "current liabilities", "流动负债合计", "流动负债总计", "流动负债总额", "流动负债"),
    "quick_assets": ("quick_assets", "速动资产", "速动资产合计"),
    "cash_and_equivalents": ("cash_and_equivalents", "货币资金"),
    "cash_equivalents_balance": ("cash_equivalents_balance", "现金及现金等价物", "现金和现金等价物", "现金及现金等价物期末余额"),
    "revenue": ("revenue", "营业收入", "营业总收入", "主营业务收入"),
    "net_profit": ("net_profit", "net profit", "净利润"),
    "gross_profit": ("gross_profit", "毛利", "毛利润"),
    "cost_of_goods_sold": ("cost_of_goods_sold", "营业成本", "主营业务成本"),
    "ebit": ("ebit", "息税前利润"),
    "interest_expense": ("interest_expense", "利息费用", "利息支出"),
    "beginning_total_assets": ("beginning_total_assets", "期初资产总计", "期初资产总额", "期初总资产"),
    "ending_total_assets": ("ending_total_assets", "期末资产总计", "期末资产总额", "期末总资产"),
    "beginning_equity": ("beginning_equity", "期初所有者权益", "期初股东权益"),
    "ending_equity": ("ending_equity", "期末所有者权益", "期末股东权益"),
    "beginning_accounts_receivable": ("beginning_accounts_receivable", "期初应收账款"),
    "ending_accounts_receivable": ("ending_accounts_receivable", "期末应收账款"),
    "beginning_inventory": ("beginning_inventory", "期初存货"),
    "ending_inventory": ("ending_inventory", "期末存货"),
    "accounts_receivable": ("accounts_receivable", "应收账款"),
    "inventory": ("inventory", "存货"),
    "operating_cash_flow": ("operating_cash_flow", "经营活动产生的现金流量净额", "经营活动现金流量净额", "经营现金流量净额"),
    "cash_flow_available_for_debt_service": ("cash_flow_available_for_debt_service", "可用于偿债现金流", "偿债现金流", "CFADS"),
    "scheduled_debt_service": ("scheduled_debt_service", "计划偿债本息", "债务本息偿付额", "现有年度偿债本息"),
    "prior_revenue": ("prior_revenue", "上期营业收入", "上年营业收入"),
    "prior_net_profit": ("prior_net_profit", "上期净利润", "上年净利润"),
    "financing_need": ("financing_need", "融资需求", "融资金额", "资金需求总额"),
    "available_internal_funds": ("available_internal_funds", "可用自有资金", "可用资金"),
}

for _key, _definition in FIELD_CATALOG.items():
    _HEADER_ALIASES[_key] = tuple(dict.fromkeys((*_HEADER_ALIASES.get(_key, ()), _key, _definition["label"])))


def _normalize_header(header: object) -> str:
    normalized = unicodedata.normalize("NFKC", str(header)).strip().casefold()
    normalized = re.sub(r"\([^)]*\)|（[^）]*）", "", normalized)
    return re.sub(r"[\s_\-—:：,，、]+", "", normalized)


_CANONICAL_HEADERS = {
    _normalize_header(alias): canonical
    for canonical, aliases in _HEADER_ALIASES.items()
    for alias in aliases
}
_MONETARY_FIELDS = frozenset(key for key, definition in FIELD_CATALOG.items() if definition["unit"] == "元")
_RATIO_FIELDS = frozenset(key for key, definition in FIELD_CATALOG.items() if definition["unit"] == "%")
_COUNT_FIELDS = frozenset(key for key, definition in FIELD_CATALOG.items() if definition["unit"] in {"年", "月"})
_NUMERIC_FIELDS = _MONETARY_FIELDS | _RATIO_FIELDS | _COUNT_FIELDS
_UNIT_MULTIPLIERS = {"元": 1, "万元": 10_000, "亿元": 100_000_000}


def _header_mapping(headers: list[object]) -> tuple[dict[object, str], dict[str, list[str]]]:
    mapping = {}
    sources: dict[str, list[str]] = {}
    for header in headers:
        canonical = _CANONICAL_HEADERS.get(_normalize_header(header))
        if canonical is None:
            continue
        mapping[header] = canonical
        sources.setdefault(canonical, []).append(str(header))
    return mapping, sources


def _find_header_row(raw: pd.DataFrame) -> int | None:
    company_candidate = None
    for row_index in range(min(len(raw), 50)):
        headers = raw.iloc[row_index].tolist()
        mapping, _ = _header_mapping(headers)
        if "company_name" not in mapping.values():
            continue
        if any(field in mapping.values() for field in FINANCIAL_COLUMNS):
            return row_index
        if company_candidate is None:
            company_candidate = row_index
    return company_candidate


def _unit_multiplier(header: object, default_unit: str | None) -> int:
    text = unicodedata.normalize("NFKC", str(header))
    for unit, multiplier in ( ("亿元", 100_000_000), ("万元", 10_000), ("元", 1) ):
        if unit in text:
            return multiplier
    return _UNIT_MULTIPLIERS.get(default_unit or "元", 1)


def _numeric_amount(value: object, multiplier: int) -> float | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, Real):
        return float(value) * multiplier
    text = str(value).strip().replace(",", "").replace("，", "")
    if text in {"", "-", "—", "–", "不适用", "N/A", "n/a"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        amount = float(text)
    except ValueError:
        return None
    return (-amount if negative else amount) * multiplier


def _numeric_ratio(value: object, header: str) -> float | None:
    """A percent sign/header means percentage points; plain canonical keys use decimals."""
    normalized_header = unicodedata.normalize("NFKC", header)
    if isinstance(value, str):
        text = unicodedata.normalize("NFKC", value).strip()
        if text.endswith("%"):
            amount = _numeric_amount(text[:-1], 1)
            return None if amount is None else amount / 100
    amount = _numeric_amount(value, 1)
    if amount is None:
        return None
    return amount / 100 if "%" in normalized_header or "百分比" in normalized_header else amount


def _infer_period(raw: pd.DataFrame, header_row: int) -> str:
    title = " ".join(
        str(value)
        for value in raw.iloc[:header_row].to_numpy().flatten()
        if value is not None and not pd.isna(value)
    )
    match = re.search(r"20\d{2}(?:年(?:第?[一二三四1-4]季度|度)?|[-/.]\d{1,2}(?:[-/.]\d{1,2})?)", title)
    return match.group(0) if match else "未注明"


def _parse_sheet(raw: pd.DataFrame, sheet_name: str) -> pd.DataFrame | None:
    # Our documentation tabs are intentionally outside the import data contract.
    if sheet_name in {"说明", "字段说明", "模拟示例"}:
        return None
    header_row = _find_header_row(raw)
    if header_row is None:
        return None
    raw_headers = raw.iloc[header_row].tolist()
    headers = [str(value).strip() if value is not None and not pd.isna(value) else f"unnamed_{index}"
               for index, value in enumerate(raw_headers)]
    mapping, source_headers = _header_mapping(headers)
    duplicate_fields = {
        canonical: names for canonical, names in source_headers.items() if len(names) > 1
    }
    if duplicate_fields:
        details = "；".join(f"{field}：{', '.join(names)}" for field, names in duplicate_fields.items())
        raise ValueError(f"多个表头对应同一字段，请每个字段只保留一列：{details}")

    frame = raw.iloc[header_row + 1:].copy()
    frame.columns = headers
    frame = frame.dropna(how="all").rename(columns=mapping)
    if "company_name" not in frame.columns:
        return None

    title = " ".join(
        str(value)
        for value in raw.iloc[:header_row].to_numpy().flatten()
        if value is not None and not pd.isna(value)
    )
    unit_match = re.search(r"(?:單位|单位)?\s*[：:]?\s*(億元|万元|萬元|元)", title)
    default_unit = unicodedata.normalize("NFKC", unit_match.group(1)) if unit_match else None
    for field in _MONETARY_FIELDS:
        if field not in frame.columns:
            continue
        source = source_headers[field][0]
        multiplier = _unit_multiplier(source, default_unit)
        frame[field] = frame[field].map(lambda value: _numeric_amount(value, multiplier))
    for field in _RATIO_FIELDS:
        if field in frame.columns:
            frame[field] = frame[field].map(lambda value: _numeric_ratio(value, source_headers[field][0]))
    for field in _COUNT_FIELDS:
        if field in frame.columns:
            frame[field] = frame[field].map(lambda value: _numeric_amount(value, 1))

    frame=frame.copy()  # Consolidate imported columns before adding report metadata.
    if "report_period" not in frame.columns:
        frame["report_period"] = _infer_period(raw, header_row)
    else:
        inferred = _infer_period(raw, header_row)
        frame["report_period"] = frame["report_period"].map(
            lambda value: inferred if value is None or pd.isna(value) or not str(value).strip() else str(value).strip()
        )
    frame["source_sheet"] = sheet_name
    frame["amount_unit"] = unicodedata.normalize("NFKC", default_unit) if default_unit else "未注明"
    return frame.reset_index(drop=True)


def _merge_same_period_sheets(frames: list[pd.DataFrame]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    mergeable: dict[tuple[str, str], dict[str, object]] = {}
    for frame in frames:
        for row in frame.to_dict(orient="records"):
            company = str(row.get("company_name", "")).strip()
            period = str(row.get("report_period", "未注明"))
            if not company or company == "nan" or company in {"请填写企业名称", "请填写公司名称"}:
                continue
            key = (company, period)
            if period == "未注明":
                rows.append(row)
                continue
            existing = mergeable.setdefault(key, {})
            for field, value in row.items():
                if field == "source_sheet":
                    sources = set(str(existing.get(field, "")).split("、")) - {""}
                    sources.add(str(value))
                    existing[field] = "、".join(sorted(sources))
                elif field == "amount_unit" and existing.get(field) not in {None, value}:
                    raise ValueError(f"同一报告期的工作表金额单位不一致：{company} {period}")
                elif field not in existing or existing[field] is None or pd.isna(existing[field]):
                    existing[field] = value
                elif value is not None and not pd.isna(value) and existing[field] != value:
                    if field in _NUMERIC_FIELDS:
                        raise ValueError(f"同一报告期工作表中的 {field} 数值冲突：{company} {period}")
            existing.setdefault("amount_unit", row.get("amount_unit", "未注明"))
    rows.extend(mergeable.values())
    return pd.DataFrame(rows)


def load_uploaded_financial_data(filename: str, content: bytes) -> pd.DataFrame:
    """Load wide tabular reports, map Chinese headers, normalize units, and retain periods."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            raw = pd.read_csv(BytesIO(content), header=None, dtype=object, encoding="utf-8-sig")
        except UnicodeDecodeError:
            raw = pd.read_csv(BytesIO(content), header=None, dtype=object, encoding="gb18030")
        raw_sheets = {"CSV": raw}
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        raw_sheets = {}
        try:
            for sheet in workbook.worksheets:
                if sheet.title in {"说明", "字段说明", "模拟示例"}:
                    continue
                rows = []
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        # Excel stores 5.2% as 0.052. Preserve that semantic before
                        # the generic table parser applies a (%) column header.
                        if isinstance(value, Real) and not isinstance(value, bool) and "%" in cell.number_format:
                            value = f"{value * 100:.15g}%"
                        values.append(value)
                    rows.append(values)
                raw_sheets[sheet.title] = pd.DataFrame(rows, dtype=object)
        finally:
            workbook.close()
    else:
        raise ValueError("仅支持 CSV 或 XLSX 格式")

    from finbridge.public_data import parse_public_blocks
    if suffix == ".csv":
        raw = raw_sheets["CSV"]
        if len(raw.columns) >= 6 and list(raw.iloc[0,:2]) == ["报表", "科目"]:
            blocks={}
            for block in raw.iloc[1:,0].dropna().unique():
                blocks[str(block)] = raw.loc[raw.iloc[:,0]==block].iloc[:,1:].reset_index(drop=True)
            raw_sheets=blocks
    public = parse_public_blocks(raw_sheets)
    if public is not None:
        return public
    parsed_sheets = [
        parsed
        for sheet_name, raw in raw_sheets.items()
        if (parsed := _parse_sheet(raw, str(sheet_name))) is not None
    ]
    if not parsed_sheets:
        raise ValueError("未找到可识别的表格：需有 company_name（公司名称）列和前 50 行内的表头")
    frame = _merge_same_period_sheets(parsed_sheets)
    for column in FINANCIAL_COLUMNS:
        if column not in frame.columns:
            frame[column] = None
    if frame.empty:
        raise ValueError("文件中没有企业记录")
    return frame
