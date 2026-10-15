"""Editable standalone Word report from the completed analysis snapshot."""
from io import BytesIO
import re


def report_docx(run):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.85)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name='Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.color.rgb=RGBColor.from_string('000000')
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.25
    for name,size in (('Title',22),('Heading 1',16),('Heading 2',12)):
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.keep_with_next=True
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Title'].font.underline=False
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('融智 Agent  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    raw=run.get('raw_public') or run.get('data') or {}
    doc.core_properties.title=str(raw.get('company_name','企业'))+'财务分析与融资决策报告'
    doc.core_properties.author='融智 Agent'
    def text(value):
        value=str(value or '')
        value=re.sub(r'(?:预计)?(?:可用于还款的?现金|还款现金)','预计可用于偿债的现金',value)
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)
    def para(value):return doc.add_paragraph(text(value))
    def heading(value):doc.add_heading(text(value),level=1)
    def cash(value):
        if value is None:return '未取得数据'
        v=value/1e8;digits=8 if 0<abs(v)<.01 else 4
        return f'{v:,.{digits}f}'.rstrip('0').rstrip('.')+'亿元'
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,h in zip(t.rows[0].cells,headers):c.text=text(h)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for c,value in zip(cells,row):c.text=text(value)
        for row in t.rows:
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    for r in p.runs:r.font.size=Pt(10)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.add_paragraph('企业财务分析与融资决策报告',style='Title')
    para(f"{raw.get('company_name','企业')}  |  {raw.get('report_period','')}  |  {raw.get('industry','')}  |  {raw.get('region','')}")
    para(f"分析编号 {run.get('id','')}  |  生成时间 {run.get('created_at','')}")
    if run.get('input_profile',{}).get('decision_scope')!='working-capital-12m':
        from finbridge.workflow import report_markdown
        for line in report_markdown(run).splitlines():
            if line.startswith('# '):continue
            if line.startswith('## '):heading(line[3:])
            elif line.strip():para(line.lstrip('- '))
    else:
        from finbridge.market_rules import effective_portfolio
        from finbridge.working_capital_year import cvar_explanation
        p=effective_portfolio(run['outputs']);f=p['funding'];risk=p['risk_evaluation'];summary=p['decision_summary']
        para('规划范围为下一完整年度的生产经营周转，期限12个月。报告汇总财务诊断、融资组合与风险检查；金额统一为亿元。')
        doc.add_heading(summary['title'],level=2);para(summary['text'])
        heading('一 财务分析')
        metrics=run['outputs'].get('calculate_ratios',{})
        def metric_value(m):
            v=m.get('value');unit=m.get('unit','')
            if v is None:return '数据不足'
            return cash(v) if unit=='元' else f'{v:.1%}' if unit=='%' else f'{v:.2f}{unit}'
        keys=['current_ratio','debt_to_asset_ratio','net_profit_margin','operating_cash_flow_to_net_profit','receivables_days','cash_conversion_cycle']
        table(['指标','本次结果'],[(metrics[k]['label'],metric_value(metrics[k])) for k in keys if k in metrics])
        for finding in run.get('financial_analysis',{}).get('year_diagnosis',[]):
            doc.add_heading(finding['title'],level=2);para(finding['text'])
        prediction=run.get('financial_analysis',{}).get('forecast',{}).get('history_prediction',{})
        future=prediction.get('future') or {};validation=prediction.get('validation') or {}
        doc.add_heading('下一年度现金预测依据',level=2)
        para(f"历史检验预测经营现金{cash(validation.get('annual_operating'))}，实际{cash(validation.get('actual'))}；有符号误差{cash(validation.get('error'))}。下一年度校正前{cash(future.get('uncalibrated_operating'))}，校正额{cash(future.get('calibration_adjustment'))}，校正后经营现金{cash(future.get('annual_operating'))}。")
        para(future.get('calibration_method','本次没有完整历史校正依据。'))
        heading('二 融资决策')
        para(f"净融资需求{cash(f['gap'])}；扣费后可用{cash(f['net_proceeds'])}；尚未覆盖{cash(f['unfunded_gap'])}。")
        selected=[c for c in p['composition'] if c['amount']>.01]
        if selected:
            table(['渠道','本金','期限','假设年利率','利息与费用'],[(c['name'],cash(c['amount']),f"{c['months']}个月",f"{c['rate']:.2%}",cash(c['total_interest']+c['total_fee'])) for c in selected])
            for c in selected:para(c['name']+'：'+c['reason'])
        else:para('本次未形成通过验证的新增融资组合；零配置不是银行拒贷或授信结论。')
        para(p.get('repayment_note',''))
        check_names={'杠杆边界','全期限滚动本息安全线','一年经营用款上限','年内支付时点敏感性'}
        rows=[]
        for check in p['verification']['checks']:
            label=check['label']
            if label not in check_names and not label.endswith('用途与资产池'):continue
            if label.endswith('用途与资产池') and '租赁' in label:continue
            actual=check['actual'];limit=check['limit']
            if label=='全期限滚动本息安全线':label='全年偿付安全余量';actual=-actual;rule='余量应不小于0'
            elif label=='年内支付时点敏感性':label='年内最小现金留存余量';rule='余量应不小于0'
            else:rule='不超过'+cash(limit)
            rows.append((label,cash(actual),rule,'通过' if check['passed'] else '未通过'))
        table(['约束','本次检验值','判定规则','结果'],rows)
        heading('三 风险评估')
        para(risk['conclusion']);para(risk['reverse_summary'])
        table(['情景','预计可用于偿债的现金','全年新旧本息','覆盖倍数','安全线缺口'],[(s['label'],cash(s['annual_cash']),cash(s['total_service']),'不适用' if s['dscr'] is None else f"{s['dscr']:.2f}倍",cash(s['gap_to_target'])) for s in risk['scenarios']])
        for s in risk['scenarios']:para(s['label']+'：'+s['rule'])
        doc.add_heading('情景CVaR的含义与用途',level=2)
        detail=risk.get('cvar_details') or cvar_explanation(risk['scenarios'],risk['cvar_alpha'])
        para(f"CVaR{risk['cvar_alpha']:.0%}为{cash(risk['cvar'])}，表示所设较差情景中企业偿债后现金余量相对基准的平均下降；包含经营现金变化和本息变化，不是贷款本金损失，也不是实际还款缺口。")
        table(['尾部情景','现金减少','本息增加','余量下降','尾部占比'],[(s['label'],cash(s['cash_decline']),cash(s['service_increase']),cash(s['loss']),f"{s['tail_share']:.2%}") for s in detail['rows']])
        para(detail['usage']);para(detail['limitations'])
        para(f"本次最紧情景达到偿债安全线还需{cash(detail['max_safety_gap'])}，应与CVaR的余量下降区别使用。")
        heading('四 方案建议与使用边界')
        for action in summary.get('actions',[]):para(action)
        para('独立验证：'+p['verification']['status']+'。'+p['verification']['note'])
        para('预计可用于偿债的现金不是实际还款支出或账面现金余额；年度现金均匀分配只用于支付时点敏感性检查，不是逐月预测。')
        ctx=run['outputs'].get('analyze_industry_region',{})
        para('行业与地区资料：'+(str(ctx.get('summary') or ctx.get('source_status') or '本次未取得有效近期市场证据。') if ctx.get('available') else '本次未取得已核验的近期行业地区资料。'))
        for source in ctx.get('sources',[])[:5]:para(str(source.get('title',''))+' '+str(source.get('url','')))
    output=BytesIO();doc.save(output);return output.getvalue()
