"""Chinese, source-oriented import workbooks and matching CSV exports."""

from __future__ import annotations

import csv
from functools import lru_cache
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from finbridge.data import FIELD_CATALOG


SIMULATED_EXAMPLE = {
    "company_name": "华辰精密制造（模拟）", "report_period": "2025年度", "industry": "高端制造",
    "region": "江苏苏州", "years_in_operation": 8, "purpose": "经营周转",
    "total_assets": 12_800_000, "total_liabilities": 5_760_000, "current_assets": 4_800_000,
    "current_liabilities": 3_200_000, "quick_assets": 3_600_000, "cash_and_equivalents": 1_000_000,
    "cash_equivalents_balance": 900_000, "restricted_cash": 100_000,
    "beginning_total_assets": 11_200_000, "ending_total_assets": 12_800_000,
    "beginning_equity": 6_500_000, "ending_equity": 7_040_000,
    "revenue": 9_600_000, "net_profit": 768_000, "gross_profit": 2_880_000,
    "cost_of_goods_sold": 6_720_000, "operating_profit": 1_060_000, "ebit": 1_200_000,
    "interest_expense": 180_000, "income_tax_expense": 252_000, "depreciation_amortization": 600_000,
    "prior_revenue": 8_400_000, "prior_net_profit": 620_000, "projected_revenue_growth": .08,
    "operating_cash_flow": 2_900_000, "cash_flow_available_for_debt_service": 2_800_000,
    "maintenance_capex": 100_000, "capital_expenditure": 300_000, "taxes_paid": 350_000,
    "scheduled_debt_service": 600_000, "annual_principal_repayment": 420_000, "annual_interest_payment": 180_000,
    "short_term_debt": 1_200_000, "long_term_debt": 2_000_000, "current_portion_long_term_debt": 300_000,
    "accounts_receivable": 2_200_000, "beginning_accounts_receivable": 1_800_000,
    "ending_accounts_receivable": 2_200_000, "receivables_overdue": 120_000,
    "receivables_over_90_days": 20_000, "eligible_receivables": 1_800_000,
    "inventory": 1_200_000, "beginning_inventory": 1_000_000, "ending_inventory": 1_200_000,
    "accounts_payable": 900_000, "beginning_accounts_payable": 700_000, "ending_accounts_payable": 900_000,
    "prepayments": 200_000, "beginning_prepayments": 180_000, "ending_prepayments": 200_000,
    "advance_receipts": 150_000, "beginning_advance_receipts": 100_000, "ending_advance_receipts": 150_000,
    "financing_need": 3_000_000, "available_internal_funds": 500_000, "confirmed_additional_internal_funds": 0, "minimum_cash_reserve": 300_000,
    "working_capital_need": 4_000_000, "working_capital_funding_need": 3_000_000,
    "equipment_funding_need": 0, "equipment_purchase_amount": 0,
    "unencumbered_collateral_value": 2_000_000, "confirmed_credit_limit": 1_000_000,
    "existing_working_capital_loans": 1_000_000, "other_working_capital_financing": 200_000,
    "credit_rate": .052, "credit_months": 24, "credit_fee_rate": .005,
    "receivable_rate": .048, "receivable_months": 12, "receivable_fee_rate": .008,
    "secured_rate": .05, "secured_months": 36, "secured_fee_rate": .01,
    "leasing_rate": .06, "leasing_months": 36, "leasing_fee_rate": .01,
    "receivable_advance_rate": .7, "collateral_advance_rate": .5, "leasing_advance_rate": .8,
}


def _headers() -> list[str]:
    return [f"{entry['label']}（{entry['unit']}）" if entry["unit"] else entry["label"] for entry in FIELD_CATALOG.values()]


def legacy_financial_template_csv(example: bool = False) -> bytes:
    """UTF-8 BOM CSV; (%) columns contain percentage-point numbers."""
    text = StringIO(newline="")
    writer = csv.writer(text)
    writer.writerow(_headers())
    values = SIMULATED_EXAMPLE if example else {"company_name": "请填写企业名称"}
    writer.writerow([
        round(values[key] * 100, 10) if key in values and entry["unit"] == "%" else values.get(key, "")
        for key, entry in FIELD_CATALOG.items()
    ])
    return text.getvalue().encode("utf-8-sig")


def _documentation(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("说明", 0)
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 27
    sheet.column_dimensions["B"].width = 105
    instructions = [
        ("企业融资分析中文模板", ""),
        ("填写主表", "只在“企业财务数据”填写。第一行是表头，每行一家企业、一个报告期。模板中的占位名称须替换；空白模板不能作为企业数据。"),
        ("金额口径", "所有带（元）的列统一填写人民币元，不用万元数值替代。若自备表格采用万元，须在列名或标题明确单位。"),
        ("利率与融资比例", "带（%）的Excel列按百分比显示，可输入5.2%或0.052，均显示5.2%；CSV同列填5.2表示5.2%。月份和经营年限不乘金额单位。"),
        ("空白和零", "未知值留空，0仅表示经核验确认为零。不要以净利润代替现金流，不要为通过检查补零。"),
        ("原始报表数据", "按同主体同期间的资产负债表、利润表、现金流量表及附注取数。各类余额、费用、实际现金收付和未来偿债计划分别填报。"),
        ("现金口径", "货币资金与现金等价物分别填报。可用自有资金须扣受限款、留存经营现金；超过已知余量的部分，另填已落实额外自有资金，排除重复现金与借款。"),
        ("偿债现金流", "可用于偿债现金流（CFADS）须由企业核验并说明口径，不等于净利润或EBITDA。已在经营现金流反映的税费不得再次机械扣除。"),
        ("本金与利息", "现有年度偿债本息是未来一年已有债务的还本付息。短期借款余额、一年内到期长期债务和全年应还本金含义不同，避免重复统计。"),
        ("融资用途与规模", "资金需求总额为本次用途预算。经营周转资金需求与设备资金需求应拆分；营运资金需求量是另一项周转测算结果，不直接当授信额度。"),
        ("渠道条件", "合格应收、担保物价值和可用信用额度须核验合同、权属、重复融资和剩余额度。利率、费率及融资比例须注明真实报价或用户假设；示例数值均为假设。"),
        ("模拟示例", "“模拟示例”是阅读用纵向对照，不导入分析。“企业融资分析示例.xlsx”的主表才包含可导入的虚构制造业案例，不代表任何真实企业或机构承诺。"),
        ("数据使用", "完成录入后上传本文件。模型规划开启时以网站提示为准发送本次企业数据；仅上传有权使用的数据。"),
    ]
    for row_index, (title, content) in enumerate(instructions, 2):
        sheet.cell(row_index, 1, title)
        sheet.cell(row_index, 2, content)
        sheet.row_dimensions[row_index].height = 42 if row_index > 2 else 28
    sheet["A2"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="243F36")
    for row in sheet.iter_rows(min_row=3):
        row[0].font = Font(name="Microsoft YaHei", size=10, bold=True, color="243F36")
        row[1].font = Font(name="Microsoft YaHei", size=10, color="37463D")
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.sheet_properties.tabColor = "8D9A87"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_area = "A2:B14"


def _reference_sheets(workbook: Workbook) -> None:
    notes = workbook.create_sheet("字段说明")
    notes.append(["字段中文名", "数据分组", "单位", "取数位置", "填报口径", "完整分析关键项"])
    sample = workbook.create_sheet("模拟示例")
    sample.append(["阅读用模拟字段", "模拟数值", "单位", "说明"])
    for key, entry in FIELD_CATALOG.items():
        # Prefixing keeps even a renamed documentation tab out of header detection.
        notes.append(["字段：" + entry["label"], entry["group"], entry["unit"] or "文本", entry["source"], entry["meaning"], "是" if entry["required"] else "按场景补充"])
        sample.append(["字段：" + entry["label"], SIMULATED_EXAMPLE.get(key), entry["unit"] or "文本", "虚构案例。渠道价格与融资比例为假设，债权、授信与担保物均设定为已由人工核验。"])
        if entry["unit"] == "%":
            sample.cell(sample.max_row, 2).number_format = "0.0%"
        elif entry["unit"] == "元":
            sample.cell(sample.max_row, 2).number_format = '#,##0.00;[Red](#,##0.00);"—"'
    for sheet, widths in ((notes, [28, 22, 12, 35, 90, 22]), (sample, [30, 32, 12, 85])):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B2"
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="344B40")
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 42
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10, color="31443A")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions


@lru_cache(maxsize=2)
def legacy_financial_template_xlsx(example: bool = False) -> bytes:
    """Return an editable workbook. Documentation/example tabs are not imported."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "企业财务数据"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "344B40"
    sheet.append(_headers())
    values = SIMULATED_EXAMPLE if example else {"company_name": "请填写企业名称"}
    sheet.append([values.get(key) for key in FIELD_CATALOG])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(FIELD_CATALOG))}31"
    group_colors = ["344B40", "4B6452", "5B6F51", "566A68", "6D6853", "4B6452", "566A68", "6D6853"]
    groups = list(dict.fromkeys(entry["group"] for entry in FIELD_CATALOG.values()))
    signed = {"net_profit", "prior_net_profit", "gross_profit", "operating_profit", "ebit", "income_tax_expense", "operating_cash_flow", "cash_flow_available_for_debt_service", "beginning_equity", "ending_equity"}
    for column, (key, entry) in enumerate(FIELD_CATALOG.items(), 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = 26 if entry["unit"] == "" else 23
        header = sheet.cell(1, column)
        header.fill = PatternFill("solid", fgColor=group_colors[groups.index(entry["group"]) % len(group_colors)])
        header.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.border = Border(right=Side(style="thin", color="FFFFFF"))
        for row in range(2, 32):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Microsoft YaHei", size=10, color="1755A5")
            cell.fill = PatternFill("solid", fgColor="FFF5D8" if entry["required"] else "F5F7F1")
            cell.alignment = Alignment(vertical="center", horizontal="left" if entry["unit"] == "" else "right")
            if entry["unit"] == "%":
                cell.number_format = "0.0%"
            elif entry["unit"] == "元":
                cell.number_format = '#,##0.00;[Red](#,##0.00);"—"'
            elif entry["unit"] in {"年", "月"}:
                cell.number_format = "0.##"
        validation = None
        if entry["unit"] == "%":
            validation = DataValidation(type="decimal", operator="between", formula1=-1 if key == "projected_revenue_growth" else 0, formula2=10 if key == "projected_revenue_growth" else 1, allow_blank=True)
            validation.promptTitle = "填写百分比"
            validation.prompt = "输入5.2%或0.052表示5.2%；未知留空。"
        elif entry["unit"] == "月":
            validation = DataValidation(type="whole", operator="between", formula1=1, formula2=120, allow_blank=True)
        elif entry["unit"] in {"元", "年"}:
            validation = DataValidation(type="decimal", operator="between", formula1=-1e15 if key in signed else 0, formula2=200 if entry["unit"] == "年" else 1e15, allow_blank=True)
        elif key == "purpose":
            validation = DataValidation(type="list", formula1='"经营周转,订单备货,设备购置,综合融资"', allow_blank=True)
        if validation is not None:
            validation.showInputMessage = True
            validation.showErrorMessage = True
            validation.errorTitle = "请核对填报口径"
            validation.error = "请填写符合单位及范围的数值，未知值留空。"
            sheet.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}1001")
    sheet.row_dimensions[1].height = 62
    for row in range(2, 32):
        sheet.row_dimensions[row].height = 24
    _documentation(workbook)
    _reference_sheets(workbook)
    workbook.active = 1  # Open at the main editable/importable table.
    result = BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


from finbridge.public_template import public_template_xlsx as financial_template_xlsx, public_template_csv as financial_template_csv
